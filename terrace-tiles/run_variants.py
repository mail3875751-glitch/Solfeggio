# -*- coding: utf-8 -*-
"""Прогон вариантов раскладки: python3 run_variants.py config/<файл>.json [out_dir]"""
import json, os, sys
from matplotlib.backends.backend_pdf import PdfPages
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill
from openpyxl.utils import get_column_letter

from tiling import Tile, build_terrace, best_layout
from render import draw_variant, write_dxf, variant_title, name_ru


def main(cfg_path, out_dir):
    cfg = json.load(open(cfg_path, encoding='utf-8'))
    os.makedirs(out_dir, exist_ok=True)
    terrace, zones = build_terrace(cfg['zones'])
    tiles = {t['name']: Tile(t['name'], t['w'], t['h']) for t in cfg['tiles']}
    joint = float(cfg['joint'])
    xc = float(cfg.get('axis_x', (terrace.bounds[0] + terrace.bounds[2]) / 2))
    yj = float(cfg['junction_y'])
    results = []
    stem = os.path.splitext(os.path.basename(cfg_path))[0]
    with PdfPages(os.path.join(out_dir, f'{stem}_варианты.pdf')) as pdf:
        for v in cfg['variants']:
            res = best_layout(terrace, zones, tiles[v['tile']], joint, v, xc, yj)
            results.append(res)
            draw_variant(res, joint, cfg['project'], cfg.get('notes', []),
                         os.path.join(out_dir, f"{stem}_{v['id']}.png"), pdf)
            write_dxf(res, joint, os.path.join(out_dir, f"{stem}_{v['id']}.dxf"))
            s = res.stats
            print(f"{variant_title(res, joint)} | целых {s['full']} | резаных {s['cut']} | "
                  f"нужно {s['tiles_with_reuse']} (без обрезков {s['tiles_no_reuse']}) | "
                  f"отход {s['waste_pct']:.1f}% | мин.кусок {s['min_piece_mm']:.0f} мм | "
                  f"<100мм: {s['slivers_lt_100']} | старт X: {s['x_mode']} | альтернативы: {s['alternatives']}")
    write_summary(cfg, results, joint, os.path.join(out_dir, f'{stem}_свод.xlsx'),
                  os.path.join(out_dir, f'{stem}_свод.md'))


def write_summary(cfg, results, joint, xlsx_path, md_path):
    heads = ['Вариант', 'Плитка, мм', 'Раскладка', 'Шов, мм', 'Площадь, м²', 'Целых, шт', 'Резаных кусков, шт',
             'Плиток без обрезков, шт', 'Плиток с обрезками, шт', 'Заказ, м² (с обрезками)', 'Отход, %',
             'Мин. кусок, мм', 'Подрезов <100 мм', 'Подрезов <⅓ плитки', 'Старт по X']
    rows = []
    for r in results:
        s, v, t = r.stats, r.variant, r.tile
        pat = {'grid': 'прямая', 'bond': 'вразбежку', 'herringbone': 'ёлочка'}[v['pattern']]
        if v['pattern'] == 'bond':
            pat += ' ½' if abs(v.get('offset', .5) - .5) < 1e-9 else ' ⅓'
        if float(v.get('angle', 0)):
            pat += f" {int(v['angle'])}°"
        rows.append([v['id'], t.name, pat, joint, round(s['area_m2'], 2), s['full'], s['cut'], s['tiles_no_reuse'],
                     s['tiles_with_reuse'], round(s['tiles_with_reuse'] * s['tile_m2'], 2), round(s['waste_pct'], 1),
                     round(s['min_piece_mm']) if s['min_piece_mm'] else None, s['slivers_lt_100'], s['slivers_lt_third'],
                     'шов на оси' if s['x_mode'] == 'joint' else 'центр плитки на оси'])
    wb = Workbook()
    ws = wb.active
    ws.title = 'Свод'
    ws['A1'] = cfg['project']
    ws['A1'].font = Font(bold=True, size=12)
    ws.append([])
    ws.append(heads)
    for c in ws[3]:
        c.font = Font(bold=True)
        c.alignment = Alignment(wrap_text=True, vertical='center', horizontal='center')
        c.fill = PatternFill('solid', fgColor='DDEBF7')
    for row in rows:
        ws.append(row)
    for i, w in enumerate([9, 11, 16, 8, 11, 10, 12, 14, 14, 14, 9, 11, 11, 12, 18], 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.row_dimensions[3].height = 45
    ws.freeze_panes = 'A4'
    r0 = ws.max_row + 2
    ws.cell(r0, 1, 'Примечания').font = Font(bold=True)
    for k, n in enumerate(cfg.get('notes', []), 1):
        ws.cell(r0 + k, 1, f'{k}. {n}')
    ws.cell(r0 + len(cfg.get('notes', [])) + 1, 1,
            'Расход «с обрезками» — расчётный минимум; к заказу добавить запас 5–7 % (прямая/вразбежку) или 10–15 % (диагональ/ёлочка) и округлить до целых коробок.')
    ws2 = wb.create_sheet('По зонам')
    ws2.append(['Вариант', 'Зона', 'Площадь, м²', 'Целых, шт', 'Резаных кусков, шт'])
    for c in ws2[1]:
        c.font = Font(bold=True)
    for r in results:
        for zn, zs in r.stats['per_zone'].items():
            ws2.append([r.variant['id'], name_ru(zn), round(zs['area_m2'], 2), zs['full'], zs['cut']])
    wb.save(xlsx_path)
    with open(md_path, 'w', encoding='utf-8') as f:
        f.write(f"# {cfg['project']}\n\n")
        f.write('| ' + ' | '.join(heads) + ' |\n|' + '---|' * len(heads) + '\n')
        for row in rows:
            f.write('| ' + ' | '.join('' if x is None else str(x) for x in row) + ' |\n')
        f.write('\n' + '\n'.join(f'- {n}' for n in cfg.get('notes', [])) + '\n')


if __name__ == '__main__':
    main(sys.argv[1], sys.argv[2] if len(sys.argv) > 2 else 'out')
